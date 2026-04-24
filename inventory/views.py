import io
import json
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.shortcuts import render, get_object_or_404, redirect
from django.http import JsonResponse, HttpResponse
from django.utils.timezone import now
from django.db import transaction
from django.db.models import F, Q, Sum
from core.access import require_store_permission
from core.views import get_store
from .models import Category, Product, ProductImage, ProductVariant, Package, PackageItem, StockMovement, Supplier, SupplierTransaction


def _find_or_create_supplier(store, name, phone=''):
    """Match supplier by name+phone; create if not found."""
    name = name.strip()
    phone = (phone or '').strip()
    qs = Supplier.objects.filter(store=store)
    if phone:
        supplier = qs.filter(name__iexact=name, phone=phone).first()
        if supplier:
            return supplier, False
    supplier = qs.filter(name__iexact=name).first()
    if supplier:
        if phone and not supplier.phone:
            supplier.phone = phone
            supplier.save(update_fields=['phone'])
        return supplier, False
    supplier = Supplier.objects.create(
        store=store, name=name, phone=phone, ownership='external',
    )
    return supplier, True


@login_required
def product_list(request, store_slug):
    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'view_inventory')
    q = request.GET.get('q', '')
    cat = request.GET.get('cat', '')
    products = Product.objects.filter(store=store).prefetch_related('variants')
    if q:
        products = products.filter(Q(name__icontains=q) | Q(sku__icontains=q))
    if cat:
        products = products.filter(category_id=cat)
    categories = Category.objects.filter(store=store)
    return render(request, 'inventory/product_list.html', {
        'store': store, 'products': products, 'categories': categories,
        'q': q, 'cat': cat,
    })


@login_required
def product_form(request, store_slug, pk=None):
    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'manage_inventory')
    product = get_object_or_404(Product, pk=pk, store=store) if pk else None
    store_categories = Category.objects.filter(store=store)
    suppliers = Supplier.objects.filter(store=store, is_active=True)

    # Backfill barcode for existing products that have none
    if product and not product.barcode:
        product.barcode = product._generate_barcode()
        product.save(update_fields=['barcode'])

    if request.method == 'POST':
        data = request.POST
        is_new = product is None
        if is_new:
            product = Product(store=store)

        product.name = data['name']
        product.sku = data.get('sku', '')
        product.description = data.get('description', '')
        product.unit_price = data['unit_price']
        product.cost_price = data.get('cost_price', 0) or 0
        product.unit_label = data.get('unit_label', 'unit')
        product.stock_qty = data.get('stock_qty', 0) or 0
        product.reorder_level = data.get('reorder_level', 5) or 5
        expiry = data.get('expiry_date', '').strip()
        product.expiry_date = expiry if expiry else None
        product.is_paid = data.get('is_paid') == 'on'

        # Category
        cat_val = data.get('category', '').strip()
        if cat_val:
            try:
                import uuid as _uuid
                _uuid.UUID(cat_val)
                product.category_id = cat_val
            except ValueError:
                cat_obj, _ = Category.objects.get_or_create(store=store, name=cat_val)
                product.category = cat_obj
        else:
            product.category = None

        # Supplier — required when stock is on credit
        sup_id = data.get('supplier', '').strip()
        if sup_id == '__new__':
            sup_id = ''
        new_sup_name = data.get('new_supplier_name', '').strip()
        new_sup_phone = data.get('new_supplier_phone', '').strip()

        if not product.is_paid:
            if sup_id:
                product.supplier_id = sup_id
            elif new_sup_name:
                sup_obj, _ = _find_or_create_supplier(store, new_sup_name, new_sup_phone)
                product.supplier = sup_obj
            else:
                return render(request, 'inventory/product_form.html', {
                    'store': store, 'product': product,
                    'store_categories': store_categories,
                    'standard_categories': [c for c in STANDARD_RETAIL_CATEGORIES
                                            if c not in set(store_categories.values_list('name', flat=True))],
                    'suppliers': suppliers,
                    'error': 'A supplier is required for stock taken on credit.',
                })
        else:
            product.supplier_id = sup_id if sup_id else None

        product.available_for_preorder = data.get('available_for_preorder') == 'on'
        lead = data.get('preorder_lead_days', '').strip()
        product.preorder_lead_days = int(lead) if lead and lead.isdigit() else None
        product.preorder_description = data.get('preorder_description', '').strip()
        product.has_color_variants = data.get('has_color_variants') == 'on'
        product.color_options = data.get('color_options', '').strip() if product.has_color_variants else ''
        product.has_size_variants = data.get('has_size_variants') == 'on'
        product.size_options = data.get('size_options', '').strip() if product.has_size_variants else ''
        product.save()

        if not product.sku:
            from .sku import generate_sku
            product.sku = generate_sku(store, product.category, product.name)
            product.save(update_fields=['sku'])

        # Record supplier debt when new product added on credit
        if is_new and not product.is_paid and product.supplier:
            cost_total = float(product.cost_price) * max(int(product.stock_qty), 1)
            if cost_total > 0:
                from django.utils.timezone import now as _now
                SupplierTransaction.objects.create(
                    supplier=product.supplier,
                    tx_type='purchase',
                    amount=cost_total,
                    description=f'Stock on credit: {product.name}',
                    date=_now().date(),
                )

        # Save variant stock when the variant form was active
        if data.get('variant_form_active') == '1':
            try:
                variants_data = json.loads(data.get('variants_json', '[]'))
            except (ValueError, TypeError):
                variants_data = []
            product.variants.all().delete()
            total = 0
            for vd in variants_data:
                size = (vd.get('size') or '').strip()
                color = (vd.get('color') or '').strip()
                qty = max(int(vd.get('qty') or 0), 0)
                ProductVariant.objects.create(product=product, size=size, color=color, stock_qty=qty)
                total += qty
            Product.objects.filter(pk=product.pk).update(stock_qty=total)

        for f in request.FILES.getlist('images'):
            ProductImage.objects.create(product=product, image=f)

        return redirect('inventory:product_list', store_slug=store.slug)

    # Build category options: existing store categories + standard ones not yet created
    existing_names = set(store_categories.values_list('name', flat=True))
    standard_not_added = [c for c in STANDARD_RETAIL_CATEGORIES if c not in existing_names]

    existing_variants_json = json.dumps({
        f"{v.size}|{v.color}": v.stock_qty
        for v in product.variants.all()
    } if product else {})

    return render(request, 'inventory/product_form.html', {
        'store': store, 'product': product,
        'store_categories': store_categories,
        'standard_categories': standard_not_added,
        'suppliers': suppliers,
        'existing_variants_json': existing_variants_json,
    })


@login_required
def stock_adjust(request, store_slug, pk):
    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'manage_inventory')
    product = get_object_or_404(Product, pk=pk, store=store)

    if request.method == 'POST':
        movement_type = request.POST['movement_type']
        qty = int(request.POST['quantity'])
        reason = request.POST.get('reason', '')

        if movement_type == 'out':
            qty = -abs(qty)
        elif movement_type == 'in':
            qty = abs(qty)

        StockMovement.objects.create(
            product=product,
            movement_type=movement_type,
            quantity=qty,
            reason=reason,
        )
        product.stock_qty = max(0, product.stock_qty + qty)
        product.save(update_fields=['stock_qty'])
        return redirect('inventory:product_list', store_slug=store.slug)

    movements = StockMovement.objects.filter(product=product)[:20]
    return render(request, 'inventory/stock_adjust.html', {
        'store': store, 'product': product, 'movements': movements,
    })


STANDARD_RETAIL_CATEGORIES = [
    'Beverages', 'Bread & Bakery', 'Breakfast & Cereal', 'Canned & Packaged Goods',
    'Condiments & Sauces', 'Confectionery & Sweets', 'Dairy & Eggs', 'Frozen Foods',
    'Fresh Produce', 'Grains & Rice', 'Meat & Poultry', 'Seafood', 'Snacks & Chips',
    'Spices & Seasonings', 'Beverages - Alcohol', 'Baby & Infant', 'Beauty & Personal Care',
    'Cleaning Supplies', 'Health & Wellness', 'Household Items', 'Laundry & Detergents',
    'Office & Stationery', 'Pet Supplies', 'Toys & Games', 'Clothing & Apparel',
    'Electronics & Accessories', 'Hardware & Tools', 'Automotive', 'Sports & Outdoor',
    'Books & Media', 'Phones & Accessories', 'Fertilizers & Agrochemicals',
    'Veterinary & Animal Feed', 'Fabric & Sewing', 'Gift & Seasonal',
]


@login_required
def category_list(request, store_slug):
    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'manage_inventory')
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        expiry_warning_days = int(request.POST.get('expiry_warning_days', 30) or 30)
        if name:
            obj, created = Category.objects.get_or_create(store=store, name=name)
            if created:
                obj.expiry_warning_days = expiry_warning_days
                obj.save(update_fields=['expiry_warning_days'])
        return redirect('inventory:category_list', store_slug=store.slug)
    categories = Category.objects.filter(store=store)
    existing_names = set(categories.values_list('name', flat=True))
    return render(request, 'inventory/category_list.html', {
        'store': store, 'categories': categories,
        'standard_categories': [c for c in STANDARD_RETAIL_CATEGORIES if c not in existing_names],
    })


@login_required
def category_form(request, store_slug):
    return category_list(request, store_slug)


@login_required
def package_list(request, store_slug):
    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'manage_inventory')
    packages = Package.objects.filter(store=store).prefetch_related('items__product')
    return render(request, 'inventory/package_list.html', {
        'store': store, 'packages': packages,
    })


@login_required
def package_form(request, store_slug, pk=None):
    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'manage_inventory')
    package = get_object_or_404(Package, pk=pk, store=store) if pk else None
    products = Product.objects.filter(store=store, is_active=True)

    if request.method == 'POST':
        data = request.POST
        if package is None:
            package = Package(store=store)
        package.name = data['name']
        package.description = data.get('description', '')
        package.package_price = data['package_price']
        package.save()

        # Clear old items and add new
        package.items.all().delete()
        product_ids = data.getlist('item_product')
        quantities = data.getlist('item_qty')
        for pid, qty in zip(product_ids, quantities):
            if pid and qty:
                PackageItem.objects.create(
                    package=package,
                    product_id=pid,
                    quantity=int(qty),
                )
        return redirect('inventory:package_list', store_slug=store.slug)

    return render(request, 'inventory/package_form.html', {
        'store': store, 'package': package, 'products': products,
    })


@login_required
def stock_overview(request, store_slug):
    from django.utils import timezone
    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'view_stock')
    products = Product.objects.filter(store=store, is_active=True).select_related('category').order_by('stock_qty')
    low = products.filter(stock_qty__lte=F('reorder_level'), stock_qty__gt=0)
    out = products.filter(stock_qty__lte=0)
    today = timezone.now().date()
    from datetime import timedelta as td
    expiring_soon = products.filter(
        expiry_date__isnull=False,
        expiry_date__gte=today,
        expiry_date__lte=today + td(days=30),
    ).order_by('expiry_date')
    expired = products.filter(expiry_date__isnull=False, expiry_date__lt=today)
    return render(request, 'inventory/stock_overview.html', {
        'store': store, 'products': products,
        'low_stock': low, 'out_of_stock': out,
        'expiring_soon': expiring_soon, 'expired': expired,
        'today': today,
    })


@login_required
def api_products(request, store_slug):
    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'view_inventory')
    q = request.GET.get('q', '')
    products = Product.objects.filter(store=store, is_active=True)
    if q:
        products = products.filter(Q(name__icontains=q) | Q(sku__icontains=q))
    data = [{
        'id': str(p.id), 'name': p.name, 'sku': p.sku,
        'unit_price': str(p.unit_price), 'cost_price': str(p.cost_price),
        'stock_qty': p.stock_qty, 'unit_label': p.unit_label,
        'category_id': str(p.category_id) if p.category_id else '',
        'category_name': p.category.name if p.category_id else '',
    } for p in products.select_related('category')[:50]]
    return JsonResponse(data, safe=False)


@login_required
def supplier_list(request, store_slug):
    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'manage_suppliers')
    suppliers = Supplier.objects.filter(store=store)
    total_owed = suppliers.filter(ownership='external').aggregate(t=Sum('outstanding_balance'))['t'] or 0
    total_invested = suppliers.filter(ownership='store_owned').aggregate(t=Sum('outstanding_balance'))['t'] or 0
    return render(request, 'inventory/supplier_list.html', {
        'store': store, 'suppliers': suppliers,
        'total_owed': total_owed, 'total_invested': total_invested,
    })


@login_required
def supplier_form(request, store_slug, pk=None):
    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'manage_suppliers')
    supplier = get_object_or_404(Supplier, pk=pk, store=store) if pk else None

    if request.method == 'POST':
        data = request.POST
        if supplier is None:
            supplier = Supplier(store=store)
        supplier.name = data['name']
        supplier.contact_person = data.get('contact_person', '')
        supplier.phone = data.get('phone', '')
        supplier.email = data.get('email', '')
        supplier.address = data.get('address', '')
        supplier.notes = data.get('notes', '')
        supplier.ownership = data.get('ownership', 'external')
        supplier.save()
        messages.success(request, f'Supplier "{supplier.name}" saved.')
        return redirect('inventory:supplier_detail', store_slug=store.slug, pk=supplier.pk)

    return render(request, 'inventory/supplier_form.html', {
        'store': store, 'supplier': supplier,
    })


@login_required
def supplier_detail(request, store_slug, pk):
    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'manage_suppliers')
    supplier = get_object_or_404(Supplier, pk=pk, store=store)
    transactions = supplier.transactions.all()[:50]
    return render(request, 'inventory/supplier_detail.html', {
        'store': store, 'supplier': supplier, 'transactions': transactions,
        'today': now().date(),
    })


@login_required
def supplier_transact(request, store_slug, pk):
    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'manage_suppliers')
    supplier = get_object_or_404(Supplier, pk=pk, store=store)

    if request.method == 'POST':
        tx_type = request.POST['tx_type']
        amount = request.POST['amount']
        description = request.POST.get('description', '')
        date = request.POST.get('date') or now().date()
        SupplierTransaction.objects.create(
            supplier=supplier, tx_type=tx_type,
            amount=amount, description=description, date=date,
        )
        messages.success(request, 'Transaction recorded.')
    return redirect('inventory:supplier_detail', store_slug=store.slug, pk=pk)


@login_required
def supplier_toggle(request, store_slug, pk):
    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'manage_suppliers')
    supplier = get_object_or_404(Supplier, pk=pk, store=store)
    supplier.is_active = not supplier.is_active
    supplier.save(update_fields=['is_active'])
    return redirect('inventory:supplier_list', store_slug=store.slug)


@login_required
def barcode_print(request, store_slug, pk):
    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'view_inventory')
    product = get_object_or_404(Product, pk=pk, store=store)
    qty = int(request.GET.get('qty', 1))
    return render(request, 'inventory/barcode_print.html', {
        'store': store, 'product': product, 'qty': range(qty),
    })


@login_required
def api_packages(request, store_slug):
    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'sell')
    packages = Package.objects.filter(store=store, is_active=True).prefetch_related('items__product')
    data = [{
        'id': str(p.id), 'name': p.name,
        'package_price': str(p.package_price),
        'items': [{'product': i.product.name, 'qty': i.quantity} for i in p.items.all()],
    } for p in packages]
    return JsonResponse(data, safe=False)


@login_required
def download_template(request, store_slug):
    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'manage_inventory')
    from .template_generator import generate_product_template
    wb = generate_product_template(store)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{store.slug}_product_template.xlsx"'
    return response


@login_required
def bulk_upload(request, store_slug):
    import os
    import re
    import zipfile
    import openpyxl
    from django.core.files.base import ContentFile
    from .sku import generate_sku

    store = get_store(store_slug, request.user)
    require_store_permission(request.user, store, 'manage_inventory')
    results = None

    if request.method == 'POST' and request.FILES.get('file'):
        created, warnings, errors = [], [], []
        images_saved = 0

        # --- Extract images from ZIP (before transaction so failures are just warnings) ---
        _ALLOWED_IMG_EXTS = {'.png', '.jpg', '.jpeg', '.webp', '.gif'}
        zip_images = {}  # stem_lower → (original_basename, file_bytes)
        if request.FILES.get('images_zip'):
            try:
                with zipfile.ZipFile(io.BytesIO(request.FILES['images_zip'].read())) as zf:
                    for zname in zf.namelist():
                        basename = os.path.basename(zname)
                        if not basename or basename.startswith('.') or basename.startswith('__'):
                            continue
                        ext = os.path.splitext(basename)[1].lower()
                        if ext not in _ALLOWED_IMG_EXTS:
                            continue
                        stem = os.path.splitext(basename)[0].lower()
                        zip_images[stem] = (basename, zf.read(zname))
                if not zip_images:
                    warnings.append('ZIP contained no recognised image files (.png .jpg .jpeg .webp .gif).')
            except zipfile.BadZipFile:
                warnings.append('Images ZIP could not be read — images not imported.')
            except Exception as zip_exc:
                warnings.append(f'Images ZIP error: {zip_exc} — images not imported.')

        # --- Load workbook ---
        try:
            wb = openpyxl.load_workbook(request.FILES['file'], data_only=True)
            ws = wb['Products'] if 'Products' in wb.sheetnames else wb.active
        except Exception as e:
            errors.append(f'Could not read file: {e}')
            results = {'created': [], 'warnings': [], 'errors': errors, 'images_saved': 0}
            return render(request, 'inventory/bulk_upload.html', {'store': store, 'results': results})

        cat_map = {c.name.lower(): c for c in Category.objects.filter(store=store)}
        product_image_queue = []  # [(product_obj, image_base_str)]

        try:
            with transaction.atomic():
                for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not any(row):
                        continue

                    name = str(row[0]).strip() if row[0] else ''
                    unit_price_raw = row[3]

                    if not name:
                        errors.append(f'Row {row_num}: Product name is required — skipped')
                        continue
                    if unit_price_raw is None or str(unit_price_raw).strip() == '':
                        errors.append(f"Row {row_num}: Unit price is required for '{name}' — skipped")
                        continue
                    try:
                        unit_price = float(unit_price_raw)
                    except (ValueError, TypeError):
                        errors.append(f"Row {row_num}: Invalid unit price '{unit_price_raw}' for '{name}' — skipped")
                        continue

                    cat_name = str(row[1]).strip().lower() if row[1] else ''
                    category = None
                    if cat_name:
                        category = cat_map.get(cat_name)
                        if not category:
                            warnings.append(f"Row {row_num}: Category '{row[1]}' not found — product created without category")

                    try:
                        stock_qty = int(row[2]) if row[2] is not None else 0
                    except (ValueError, TypeError):
                        stock_qty = 0

                    cost_price = 0
                    try:
                        if row[4] is not None:
                            cost_price = float(row[4])
                    except (ValueError, TypeError):
                        pass

                    unit_label = str(row[5]).strip() if row[5] else 'unit'
                    supplier_name = str(row[6]).strip() if row[6] else ''
                    supplier_phone = str(row[7]).strip() if row[7] else ''
                    paid_val = str(row[8]).strip().upper() if row[8] else 'YES'
                    is_paid = paid_val != 'NO'
                    preorder_val = str(row[9]).strip().upper() if row[9] else 'NO'
                    available_for_preorder = preorder_val == 'YES'
                    try:
                        reorder_level = int(row[10]) if row[10] is not None else 5
                    except (ValueError, TypeError):
                        reorder_level = 5

                    # New columns: size options (L), color options (M), image filename (N)
                    size_opts_raw = str(row[11]).strip() if len(row) > 11 and row[11] else ''
                    color_opts_raw = str(row[12]).strip() if len(row) > 12 and row[12] else ''
                    image_base = ''
                    if len(row) > 13 and row[13]:
                        # Strip extension if user included it (e.g. "sneakers.png" → "sneakers")
                        image_base = os.path.splitext(str(row[13]).strip())[0].strip()

                    if not is_paid and not supplier_name:
                        errors.append(f"Row {row_num}: '{name}' is marked unpaid — Supplier Name is required. Skipped.")
                        continue

                    supplier_obj = None
                    if supplier_name:
                        supplier_obj, _ = _find_or_create_supplier(store, supplier_name, supplier_phone)

                    product = Product(
                        store=store,
                        category=category,
                        supplier=supplier_obj,
                        name=name,
                        unit_price=unit_price,
                        cost_price=cost_price,
                        unit_label=unit_label,
                        stock_qty=stock_qty,
                        reorder_level=reorder_level,
                        is_paid=is_paid,
                        available_for_preorder=available_for_preorder,
                        has_size_variants=bool(size_opts_raw),
                        size_options=size_opts_raw,
                        has_color_variants=bool(color_opts_raw),
                        color_options=color_opts_raw,
                    )
                    product.save()
                    product.sku = generate_sku(store, category, name)
                    product.save(update_fields=['sku'])

                    if stock_qty > 0:
                        StockMovement.objects.create(
                            product=product, movement_type='in',
                            quantity=stock_qty, reason='Bulk upload',
                        )

                    if not is_paid and supplier_obj:
                        cost_total = cost_price * max(stock_qty, 1)
                        if cost_total > 0:
                            from django.utils import timezone as _tz
                            SupplierTransaction.objects.create(
                                supplier=supplier_obj,
                                tx_type='purchase',
                                amount=cost_total,
                                description=f'Stock on credit (bulk upload): {name}',
                                date=_tz.now().date(),
                            )

                    created.append(name)
                    if image_base:
                        product_image_queue.append((product, image_base))

        except Exception as e:
            errors.append(f'Upload failed: {e}')
            product_image_queue.clear()  # products rolled back — nothing to attach images to

        # --- Process Variants sheet (outside main transaction, fail-safe) ---
        if 'Variants' in wb.sheetnames and created:
            ws_var = wb['Variants']
            variant_product_ids = set()
            for vrow_num, vrow in enumerate(ws_var.iter_rows(min_row=2, values_only=True), start=2):
                if not any(vrow):
                    continue
                prod_name = str(vrow[0]).strip() if vrow[0] else ''
                size = str(vrow[1]).strip() if len(vrow) > 1 and vrow[1] else ''
                color = str(vrow[2]).strip() if len(vrow) > 2 and vrow[2] else ''
                try:
                    qty = max(int(vrow[3] or 0), 0) if len(vrow) > 3 else 0
                except (TypeError, ValueError):
                    qty = 0
                if not prod_name:
                    continue
                try:
                    prod_obj = Product.objects.get(store=store, name__iexact=prod_name, is_active=True)
                except Product.DoesNotExist:
                    warnings.append(f"Variants row {vrow_num}: '{prod_name}' not found in store — skipped")
                    continue
                except Product.MultipleObjectsReturned:
                    prod_obj = Product.objects.filter(store=store, name__iexact=prod_name, is_active=True).first()

                ProductVariant.objects.update_or_create(
                    product=prod_obj, size=size, color=color,
                    defaults={'stock_qty': qty},
                )
                # Mark product options if not already set
                update_fields = []
                if size and not prod_obj.has_size_variants:
                    prod_obj.has_size_variants = True
                    existing_sizes = set(prod_obj.size_option_list)
                    existing_sizes.add(size)
                    prod_obj.size_options = ', '.join(sorted(existing_sizes))
                    update_fields += ['has_size_variants', 'size_options']
                if color and not prod_obj.has_color_variants:
                    prod_obj.has_color_variants = True
                    existing_colors = set(prod_obj.color_option_list)
                    existing_colors.add(color)
                    prod_obj.color_options = ', '.join(sorted(existing_colors))
                    update_fields += ['has_color_variants', 'color_options']
                if update_fields:
                    prod_obj.save(update_fields=update_fields)
                variant_product_ids.add(prod_obj.pk)

            # Sync each affected product's total stock_qty from its variants
            for pid in variant_product_ids:
                total = ProductVariant.objects.filter(product_id=pid).aggregate(t=Sum('stock_qty'))['t'] or 0
                Product.objects.filter(pk=pid).update(stock_qty=total)

        # --- Match and save images outside the transaction (fail-safe per product) ---
        if product_image_queue and zip_images:
            for product, img_base in product_image_queue:
                base_lower = img_base.lower()
                # Match exact base OR base followed by digits: productA, productA1, productA2, …
                pattern = re.compile(r'^' + re.escape(base_lower) + r'\d*$')
                matches = sorted(
                    ((stem, fname, data) for stem, (fname, data) in zip_images.items() if pattern.match(stem)),
                    key=lambda x: x[0],
                )
                if not matches:
                    warnings.append(f"'{product.name}': no image matching '{img_base}' found in ZIP — skipped")
                    continue
                for sort_order, (_, orig_fname, img_bytes) in enumerate(matches):
                    try:
                        ext = os.path.splitext(orig_fname)[1] or '.jpg'
                        pi = ProductImage(product=product, sort_order=sort_order)
                        pi.image.save(f'{product.pk}_{sort_order}{ext}', ContentFile(img_bytes), save=True)
                        images_saved += 1
                    except Exception as img_err:
                        warnings.append(f"Image '{orig_fname}' for '{product.name}': {img_err}")
        elif product_image_queue and not zip_images:
            warnings.append(
                f'{len(product_image_queue)} product(s) have image filenames specified but no ZIP was uploaded — '
                'upload an images ZIP to attach product photos.'
            )

        results = {'created': created, 'warnings': warnings, 'errors': errors, 'images_saved': images_saved}

    return render(request, 'inventory/bulk_upload.html', {'store': store, 'results': results})
