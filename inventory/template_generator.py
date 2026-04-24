from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import quote_sheetname


HEADERS = [
    ('Product Name *', 30),
    ('Category', 20),
    ('Quantity', 12),
    ('Unit Price *', 14),
    ('Cost Price', 14),
    ('Unit Label', 14),
    ('Supplier Name', 22),
    ('Supplier Phone', 18),
    ('Paid (YES/NO)', 14),
    ('Preorder (YES/NO)', 18),
    ('Reorder Level', 14),
    ('Size Options', 22),
    ('Color Options', 22),
    ('Image Filename', 22),
]


def _opts_to_csv(text):
    """Normalise newline-or-comma separated options to a single comma-separated string."""
    if not text:
        return ''
    import re
    parts = [p.strip() for p in re.split(r'[,\n]+', text) if p.strip()]
    return ', '.join(parts)


def generate_product_template(store):
    wb = Workbook()

    header_fill = PatternFill('solid', fgColor='D8F3DC')
    existing_fill = PatternFill('solid', fgColor='EBF5FB')  # light blue for existing rows
    header_font = Font(bold=True)

    # ── Products sheet ───────────────────────────────────────────────────
    ws = wb.active
    ws.title = 'Products'

    for col, (header, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = 'A2'

    # YES/NO validation for Paid (col 9) and Preorder (col 10)
    yn_dv = DataValidation(type='list', formula1='"YES,NO"', allow_blank=True)
    yn_dv.sqref = 'I2:J1048576'
    ws.add_data_validation(yn_dv)

    # Comments
    ws['A1'].comment = Comment(
        'Existing products are pre-filled below (blue rows).\n'
        'Add new products beneath them.\n'
        'Product Name and Unit Price are required.\nSKU will be auto-generated.',
        'SalesApp',
    )
    ws['H1'].comment = Comment(
        'If Paid = NO, Supplier Name is required.\nPhone helps identify existing suppliers.',
        'SalesApp',
    )
    ws['L1'].comment = Comment(
        'Comma-separated size variants, e.g: S, M, L, XL\nLeave blank if product has no sizes.',
        'SalesApp',
    )
    ws['M1'].comment = Comment(
        'Comma-separated colour variants, e.g: Red, Blue, Green\nLeave blank if product has no colours.',
        'SalesApp',
    )
    ws['N1'].comment = Comment(
        'Base filename of the product image WITHOUT extension.\n'
        'e.g. sneakers_red   →  matches sneakers_red.png\n'
        'Multiple images: sneakers_red1.png, sneakers_red2.png\n'
        'Upload a ZIP of all images in the form alongside this sheet.',
        'SalesApp',
    )

    # Pre-fill existing products
    products = (
        store.products
        .select_related('category', 'supplier')
        .prefetch_related('variants')
        .order_by('name')
    )
    row = 2
    for p in products:
        data = [
            p.name,
            p.category.name if p.category else '',
            p.stock_qty,
            float(p.unit_price),
            float(p.cost_price) if p.cost_price else '',
            p.unit_label,
            p.supplier.name if p.supplier else '',
            p.supplier.phone if p.supplier and p.supplier.phone else '',
            'YES' if p.is_paid else 'NO',
            'YES' if p.available_for_preorder else 'NO',
            p.reorder_level,
            _opts_to_csv(p.size_options),
            _opts_to_csv(p.color_options),
            '',  # image filename — not stored, left blank
        ]
        for col, val in enumerate(data, start=1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.fill = existing_fill
        row += 1

    # ── Categories reference sheet ───────────────────────────────────────
    # Include this store's categories first, then any unique names from other stores.
    from inventory.models import Category as _Category
    store_cats = list(store.categories.order_by('name').values_list('name', flat=True))
    seen = {n.lower() for n in store_cats}
    other_cats = (
        _Category.objects
        .exclude(store=store)
        .order_by('name')
        .values_list('name', flat=True)
    )
    extra = []
    for name in other_cats:
        if name.lower() not in seen:
            seen.add(name.lower())
            extra.append(name)
    all_cat_names = list(store_cats) + sorted(extra)

    ws2 = wb.create_sheet('Categories')
    hdr = ws2.cell(row=1, column=1, value='Category Name')
    hdr.font = header_font
    hdr.fill = header_fill
    ws2.column_dimensions['A'].width = 28
    for i, name in enumerate(all_cat_names, start=2):
        ws2.cell(row=i, column=1, value=name)

    ws2.sheet_state = 'visible'

    # Named range so the Products sheet can reference it as a dropdown
    cat_count = max(len(all_cat_names), 1)
    cat_range = f"{quote_sheetname('Categories')}!$A$2:$A${cat_count + 1}"
    wb.defined_names['CategoryList'] = DefinedName('CategoryList', attr_text=cat_range)

    # Dropdown validation on Category column (B) in Products sheet
    cat_dv = DataValidation(type='list', formula1='CategoryList', allow_blank=True, showErrorMessage=False)
    cat_dv.sqref = 'B2:B1048576'
    ws.add_data_validation(cat_dv)

    # ── Variants sheet ───────────────────────────────────────────────────
    ws3 = wb.create_sheet('Variants')
    var_headers = [
        ('Product Name', 30),
        ('Size', 15),
        ('Color', 15),
        ('Quantity', 14),
    ]
    for col, (h, w) in enumerate(var_headers, start=1):
        cell = ws3.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        ws3.column_dimensions[get_column_letter(col)].width = w

    ws3.freeze_panes = 'A2'
    ws3['A1'].comment = Comment(
        'Products with size/colour variations are pre-filled here (blue rows).\n'
        'Each row is one variant combination with its stock quantity.\n'
        'Product Name must match exactly a product in the Products sheet.\n'
        'Leave Size or Color blank if the product only varies on one dimension.\n'
        'The Quantity in the Products sheet is ignored for products listed here —\n'
        'total stock is calculated as the sum of all variant rows.',
        'SalesApp',
    )

    # Pre-fill existing variants (only products that actually have variants)
    vrow = 2
    for p in products:
        variants = list(p.variants.all())
        if not variants:
            continue
        for v in variants:
            vdata = [p.name, v.size, v.color, v.stock_qty]
            for col, val in enumerate(vdata, start=1):
                cell = ws3.cell(row=vrow, column=col, value=val)
                cell.fill = existing_fill
            vrow += 1

    ws3.sheet_state = 'visible'

    return wb
